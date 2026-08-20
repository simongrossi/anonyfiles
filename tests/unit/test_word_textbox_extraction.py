"""Zones de texte .docx — issue #78.

python-docx n'expose que les ``w:p`` enfants directs du corps, des cellules de
tableau et des en-têtes/pieds. Le texte d'une zone de texte vit dans un
``w:txbxContent`` imbriqué : il échappait entièrement à l'anonymisation.
"""

import pytest

pytest.importorskip("docx")

from docx import Document
from docx.oxml.ns import nsmap, qn
from docx.oxml.parser import parse_xml

from anonyfiles_core.anonymizer.word_processor import DocxProcessor

_NS = " ".join(f'xmlns:{prefix}="{uri}"' for prefix, uri in nsmap.items())

# Zone de texte VML (`w:pict`), la forme héritée toujours produite par Word.
VML_TEXTBOX = f"""
<w:p {_NS} xmlns:v="urn:schemas-microsoft-com:vml">
  <w:r>
    <w:pict>
      <v:shape style="width:200pt;height:50pt">
        <v:textbox>
          <w:txbxContent>
            <w:p><w:r><w:t>{{text}}</w:t></w:r></w:p>
          </w:txbxContent>
        </v:textbox>
      </v:shape>
    </w:pict>
  </w:r>
</w:p>
"""

# Zone de texte DrawingML enveloppée dans un `mc:AlternateContent` : Word y
# duplique le contenu (Choice moderne + Fallback VML). Les DEUX copies doivent
# être anonymisées, sinon le repli conserve le texte d'origine en clair.
DRAWINGML_TEXTBOX = f"""
<w:p {_NS}
     xmlns:v="urn:schemas-microsoft-com:vml"
     xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"
     xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
  <w:r>
    <mc:AlternateContent>
      <mc:Choice Requires="wps">
        <w:drawing>
          <wp:inline xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing">
            <a:graphic xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
              <a:graphicData>
                <wps:wsp>
                  <wps:txbx>
                    <w:txbxContent>
                      <w:p><w:r><w:t>{{text}}</w:t></w:r></w:p>
                    </w:txbxContent>
                  </wps:txbx>
                </wps:wsp>
              </a:graphicData>
            </a:graphic>
          </wp:inline>
        </w:drawing>
      </mc:Choice>
      <mc:Fallback>
        <w:pict>
          <v:shape style="width:200pt;height:50pt">
            <v:textbox>
              <w:txbxContent>
                <w:p><w:r><w:t>{{text}}</w:t></w:r></w:p>
              </w:txbxContent>
            </v:textbox>
          </v:shape>
        </w:pict>
      </mc:Fallback>
    </mc:AlternateContent>
  </w:r>
</w:p>
"""


def _append_xml(container_element, xml: str, text: str) -> None:
    container_element.append(parse_xml(xml.format(text=text).strip()))


def _build_docx(tmp_path, *, header_textbox: bool = False):
    doc = Document()
    doc.add_paragraph("Le corps mentionne Jean Dupont.")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).paragraphs[0].text = "La cellule mentionne Marie Curie."
    _append_xml(doc.element.body, VML_TEXTBOX, "La zone VML mentionne Paul Valery.")
    _append_xml(
        doc.element.body, DRAWINGML_TEXTBOX, "La zone DrawingML mentionne Alan Turing."
    )
    if header_textbox:
        header = doc.sections[0].header
        header.is_linked_to_previous = False
        _append_xml(header._element, VML_TEXTBOX, "L'en-tete mentionne Ada Lovelace.")

    path = tmp_path / "avec_zones_de_texte.docx"
    doc.save(path)
    return path


def test_textbox_text_is_extracted(tmp_path):
    path = _build_docx(tmp_path)
    blocks = DocxProcessor().extract_blocks(path)

    assert "La zone VML mentionne Paul Valery." in blocks
    # Choice + Fallback : le contenu apparaît deux fois, et les deux doivent
    # être traités pour ne pas laisser de texte en clair dans le repli.
    assert blocks.count("La zone DrawingML mentionne Alan Turing.") == 2
    # Le contenu déjà géré n'est pas dupliqué par le nouveau parcours.
    assert blocks.count("Le corps mentionne Jean Dupont.") == 1
    assert blocks.count("La cellule mentionne Marie Curie.") == 1


def test_textbox_in_header_is_extracted(tmp_path):
    path = _build_docx(tmp_path, header_textbox=True)
    blocks = DocxProcessor().extract_blocks(path)

    assert "L'en-tete mentionne Ada Lovelace." in blocks


def test_textbox_text_is_rewritten_in_output(tmp_path):
    """Le tour complet : extraction -> substitution -> réécriture du .docx."""
    path = _build_docx(tmp_path)
    processor = DocxProcessor()
    blocks = processor.extract_blocks(path)

    anonymised = [block.replace("Alan Turing", "PER_1") for block in blocks]
    anonymised = [block.replace("Paul Valery", "PER_2") for block in anonymised]

    output = tmp_path / "sortie.docx"
    processor.reconstruct_and_write_anonymized_file(output, anonymised, path)

    reread = processor.extract_blocks(output)
    assert "La zone VML mentionne PER_2." in reread
    assert reread.count("La zone DrawingML mentionne PER_1.") == 2

    # Aucune trace du texte d'origine où que ce soit dans le XML livré.
    rewritten = Document(str(output))
    full_xml = rewritten.element.xml
    assert "Alan Turing" not in full_xml
    assert "Paul Valery" not in full_xml


def test_document_without_textbox_is_unaffected(tmp_path):
    doc = Document()
    doc.add_paragraph("Un simple paragraphe.")
    path = tmp_path / "simple.docx"
    doc.save(path)

    assert DocxProcessor().extract_blocks(path) == ["Un simple paragraphe."]


def test_paragraph_lookup_ignores_non_textbox_nesting(tmp_path):
    """Un `w:p` du corps n'est jamais confondu avec du contenu de zone de texte."""
    doc = Document()
    paragraph = doc.add_paragraph("Corps.")
    assert DocxProcessor._is_inside_textbox(paragraph._p) is False

    _append_xml(doc.element.body, VML_TEXTBOX, "Dans la zone.")
    textbox_p = [
        p
        for p in doc.element.body.iter(qn("w:p"))
        if DocxProcessor._is_inside_textbox(p)
    ]
    assert len(textbox_p) == 1
