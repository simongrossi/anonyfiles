from anonymizer.word_processor import DocxProcessor
from docx import Document
import tempfile, os

def test_extract_blocks_docx():
    tmp = tempfile.NamedTemporaryFile("w+b", delete=False, suffix=".docx")
    doc = Document()
    doc.add_paragraph("Pierre est à Paris.")
    doc.add_paragraph("Ceci est un test.")
    doc.save(tmp.name)
    processor = DocxProcessor()
    blocks = processor.extract_blocks(tmp.name)
    assert len(blocks) == 2
    assert "Pierre" in blocks[0]

def test_replace_entities_docx():
    tmp_in = tempfile.NamedTemporaryFile("w+b", delete=False, suffix=".docx")
    tmp_out = tempfile.NamedTemporaryFile("w+b", delete=False, suffix=".docx")
    doc = Document()
    doc.add_paragraph("Pierre à Paris.")
    doc.save(tmp_in.name)
    processor = DocxProcessor()
    replacements = {"Pierre": "NOM003", "Paris": "VILLE_Z"}
    entities_offsets = [[("Pierre", "PER", 0, 6), ("Paris", "LOC", 9, 14)]]
    processor.replace_entities(tmp_in.name, tmp_out.name, replacements, entities_offsets)
    doc_result = Document(tmp_out.name)
    assert "NOM003" in doc_result.paragraphs[0].text
    assert "VILLE_Z" in doc_result.paragraphs[0].text
